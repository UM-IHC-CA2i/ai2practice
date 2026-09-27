#!/usr/bin/env python3
"""Build the downloadable FDA AI-Enabled Medical Device List PowerPoint slide.

Uses python-pptx so this works locally and in GitHub Actions without Node/npm.
Run after scripts/update_fda_ai_data.py has created data/fda-ai-summary.json.
"""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
import math
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data" / "fda-ai-summary.json"
OUT_DIR = ROOT / "downloads"
OUT_PATH = OUT_DIR / "fda-ai-enabled-medical-device-list-slide.pptx"
SITE_URL = "https://um-ihc-ca2i.github.io/ai2practice/tracker.html"
AUTHORS = "Li S, Andriole K, Chang P, Grimm L, Heilbrun M, Krupinski E, Klontzas M, Amiruddin R, Kadom N, Doo F"
COLORS = [
    "2568B8", "C93636", "E97817", "16858C", "AD2458", "F3A62A", "3E8B45",
    "6A2CA0", "4B2CA3", "14725E", "8B7468", "7F97A1", "D44A1C", "5A8D3B", "64748B",
]


def hex_to_rgb(value: str) -> RGBColor:
    value = value.strip().lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def fmt(n: Any) -> str:
    try:
        return f"{int(n):,}"
    except Exception:
        return "—"


def pct(n: Any) -> str:
    try:
        val = float(n)
        return f"{val:.1f}".rstrip("0").rstrip(".") + "%"
    except Exception:
        return "—"


def date_label(value: str | None) -> str:
    if not value:
        return "not available"
    value = str(value)
    try:
        if len(value) == 10:
            dt = datetime.strptime(value, "%Y-%m-%d")
        else:
            dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return dt.strftime("%b %-d, %Y")
    except Exception:
        # Windows does not support %-d; use a safer fallback.
        try:
            if len(value) == 10:
                dt = datetime.strptime(value, "%Y-%m-%d")
            else:
                dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
            return dt.strftime("%b %d, %Y").replace(" 0", " ")
        except Exception:
            return value


def add_textbox(slide, text: str, x: float, y: float, w: float, h: float, *,
                size: float = 12, bold: bool = False, italic: bool = False,
                color: str = "10233F", align=PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = hex_to_rgb(color)


def add_metric(slide, x: float, y: float, value: str, label: str) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.08), Inches(0.86))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = hex_to_rgb("DDE6EC")
    shape.line.width = Pt(0.75)
    add_textbox(slide, value, x + 0.08, y + 0.13, 0.92, 0.26, size=17, bold=True, color="10233F", align=PP_ALIGN.CENTER)
    add_textbox(slide, label, x + 0.06, y + 0.48, 0.96, 0.25, size=6.8, color="64748B", align=PP_ALIGN.CENTER)


def nice_max(value: int) -> int:
    if value <= 10:
        return 10
    magnitude = 10 ** (len(str(value)) - 1)
    normalized = value / magnitude
    if normalized <= 2:
        step = 0.5
    elif normalized <= 5:
        step = 1
    else:
        step = 2
    return int(math.ceil(normalized / step) * step * magnitude)


def main() -> None:
    if not DATA_PATH.exists():
        raise SystemExit("Missing data/fda-ai-summary.json. Run: python scripts/update_fda_ai_data.py")

    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    years = [y for y in data.get("years", []) if int(y.get("total", 0) or 0) > 0]
    if not years:
        raise SystemExit("No yearly data available in data/fda-ai-summary.json.")

    panels = list(data.get("display_panels") or [])
    if not panels:
        panels = ["Total"]

    labels = [str(y.get("year")) for y in years]
    chart_data = CategoryChartData()
    chart_data.categories = labels
    for panel in panels:
        values = []
        for y in years:
            panel_values = y.get("panels") or {}
            if panel in panel_values:
                values.append(int(panel_values.get(panel) or 0))
            elif panel == "Total":
                values.append(int(y.get("total") or 0))
            else:
                values.append(0)
        chart_data.add_series(panel, values)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = hex_to_rgb("FBF8F0")

    add_textbox(slide, "FDA AI-Enabled Medical Device List", 0.55, 0.32, 7.6, 0.38, size=22, bold=True)
    add_textbox(slide, "Entries by final-decision year and FDA lead panel", 0.55, 0.75, 7.7, 0.24, size=11, color="64748B")
    add_textbox(slide, "Source data: U.S. Food and Drug Administration AI-Enabled Medical Devices List CSV", 0.55, 1.03, 8.4, 0.23, size=8.5, color="64748B")

    add_metric(slide, 9.02, 0.28, fmt(data.get("total_list_entries")), "FDA list entries")
    add_metric(slide, 10.38, 0.28, fmt(data.get("radiology_entries")), "Radiology entries")
    add_metric(slide, 11.75, 0.28, pct(data.get("radiology_share_pct")), "Radiology share")

    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED,
        Inches(0.55), Inches(1.42), Inches(10.05), Inches(4.55),
        chart_data,
    )
    chart = chart_shape.chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.font.size = Pt(8)
    chart.category_axis.tick_labels.font.size = Pt(7)
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    chart.value_axis.maximum_scale = nice_max(max(int(y.get("total") or 0) for y in years))
    chart.value_axis.minimum_scale = 0
    try:
        chart.value_axis.axis_title.text_frame.text = "Number of devices"
        chart.value_axis.has_title = True
    except Exception:
        pass

    for i, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = hex_to_rgb(COLORS[i % len(COLORS)])
        series.format.line.color.rgb = hex_to_rgb(COLORS[i % len(COLORS)])

    sync_date = date_label(str(data.get("synced_at_utc", ""))[:10])
    latest_date = date_label(data.get("latest_final_decision"))
    add_textbox(slide, f"Latest final-decision date in FDA source: {latest_date}   |   AI2Practice sync: {sync_date}", 0.55, 6.04, 12.15, 0.18, size=7.5, color="64748B")

    citation = (
        f"{AUTHORS}. AI2Practice: FDA AI-Enabled Medical Device List visualization. "
        f"AI2Practice. {SITE_URL}. Source data: U.S. FDA AI-Enabled Medical Devices List. "
        "Visualization adapted by AI2Practice for educational use; not an FDA product or endorsement."
    )
    add_textbox(slide, citation, 0.55, 6.32, 12.20, 0.55, size=6.5, color="334155")
    add_textbox(slide, "FDA notes that the AI-Enabled Medical Device List is updated periodically and is not comprehensive.", 0.55, 6.94, 12.20, 0.18, size=6.5, italic=True, color="64748B")

    # Store source/citation details in the presentation metadata as a practical equivalent to notes.
    props = prs.core_properties
    props.author = "AI2Practice"
    props.title = "FDA AI-Enabled Medical Device List"
    props.subject = "FDA AI-Enabled Medical Device List visualization"
    props.keywords = "FDA; AI-enabled medical devices; AI2Practice; radiology"
    props.comments = "Source: FDA AI-Enabled Medical Devices List; visualization adapted by AI2Practice for educational use."
    props.modified = datetime.now(timezone.utc)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PATH)
    print(f"Wrote: {OUT_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
